"""Word 打分表生成（F8/BR-4）：python-docx OOXML 直写。

优先使用 TEMPLATE_WORD_PATH 模板填充；模板缺失或损坏时降级为内置排版。
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.settings import get_settings

_STAGE_LABEL = {"summary": "总结评审", "proposal": "立项评审"}


def _fill_table(table, template: dict[str, Any], opinions: dict[str, str]) -> None:
    """向已有表格填充维度行（维度/权重/评分区间/专家评分/评审意见）。"""
    for dim in template["dimensions"]:
        row = table.add_row()
        values = [
            dim["name"],
            str(dim["weight"]),
            f"{int(dim['weight'] * 0.75)}~{dim['weight']}",
            "",
            opinions.get(dim["name"], ""),
        ]
        for cell, value in zip(row.cells, values):
            cell.text = value


def _build_builtin(
    out_path: Path,
    *,
    project_name: str,
    stage: str,
    template: dict[str, Any],
    opinions: dict[str, str],
) -> None:
    """内置排版降级：合并单元格表头 + 维度行。"""
    doc = Document()
    title = doc.add_heading(f"项目评审打分表（{_STAGE_LABEL.get(stage, stage)}）", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    table = doc.add_table(rows=2, cols=5)
    table.style = "Table Grid"

    # 表头第 1 行：项目名称标签 + 合并单元格填项目名
    head = table.rows[0]
    head.cells[0].text = "项目名称"
    merged = head.cells[1].merge(head.cells[2]).merge(head.cells[3]).merge(head.cells[4])
    merged.text = project_name

    # 表头第 2 行：列名
    columns = ["评审维度", "权重", "评分区间", "专家评分", "评审意见"]
    for cell, name in zip(table.rows[1].cells, columns):
        cell.text = name

    _fill_table(table, template, opinions)
    doc.save(str(out_path))


def _fill_from_template(
    out_path: Path,
    *,
    project_name: str,
    stage: str,
    template: dict[str, Any],
    opinions: dict[str, str],
) -> None:
    """基于既有模板 docx 填充：向第一个表格追加维度行，并写入标题信息。"""
    doc = Document(str(get_settings().template_word_file))
    doc.add_paragraph(f"项目名称：{project_name}（{_STAGE_LABEL.get(stage, stage)}）")
    if doc.tables:
        _fill_table(doc.tables[0], template, opinions)
    else:
        # 模板无表格则退化为内置排版追加
        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for cell, name in zip(
            table.rows[0].cells, ["评审维度", "权重", "评分区间", "专家评分", "评审意见"]
        ):
            cell.text = name
        _fill_table(table, template, opinions)
    doc.save(str(out_path))


def build_score_sheet(
    *,
    task_id: str,
    project_name: str,
    stage: str,
    template: dict[str, Any],
    opinions: dict[str, str],
    version: int,
) -> Path:
    """生成打分表 docx，返回产物路径。模板缺失/异常自动降级内置排版。"""
    out_dir = get_settings().artifacts_dir_path / task_id
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"score_sheet_{stage}_v{version}.docx"

    template_file = get_settings().template_word_file
    if template_file.exists():
        try:
            _fill_from_template(
                out_path,
                project_name=project_name,
                stage=stage,
                template=template,
                opinions=opinions,
            )
            return out_path
        except Exception:
            pass  # 降级内置排版
    _build_builtin(
        out_path,
        project_name=project_name,
        stage=stage,
        template=template,
        opinions=opinions,
    )
    return out_path
