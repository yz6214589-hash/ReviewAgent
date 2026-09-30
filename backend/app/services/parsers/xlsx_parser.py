from pathlib import Path

from openpyxl import load_workbook

from app.services.parsers.base import BaseParser, ParsedDocument


class XlsxParser(BaseParser):
    file_type = "xlsx"

    def parse(self, path: Path) -> ParsedDocument:
        wb = load_workbook(str(path), read_only=True, data_only=True)
        parts: list[str] = []
        for sheet in wb.worksheets:
            parts.append(f"# Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    parts.append(" | ".join(cells))
        wb.close()
        return ParsedDocument(
            filename=path.name,
            file_type=self.file_type,
            text="\n".join(parts),
            meta={"sheets": len(wb.sheetnames)},
        )
