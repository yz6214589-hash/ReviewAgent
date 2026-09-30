from pathlib import Path

from pptx import Presentation

from app.services.parsers.base import BaseParser, ParsedDocument


class PptxParser(BaseParser):
    file_type = "pptx"

    def parse(self, path: Path) -> ParsedDocument:
        prs = Presentation(str(path))
        parts: list[str] = []
        for idx, slide in enumerate(prs.slides, start=1):
            parts.append(f"# Slide {idx}")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text:
                        parts.append(text)
        return ParsedDocument(
            filename=path.name,
            file_type=self.file_type,
            text="\n".join(parts),
            meta={"slides": len(prs.slides._sldIdLst)},
        )
