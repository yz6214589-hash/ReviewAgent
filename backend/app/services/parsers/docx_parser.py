from pathlib import Path

from docx import Document

from app.services.parsers.base import BaseParser, ParsedDocument


class DocxParser(BaseParser):
    file_type = "docx"

    def parse(self, path: Path) -> ParsedDocument:
        doc = Document(str(path))
        parts: list[str] = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                parts.append(text)
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                line = " | ".join([c for c in cells if c])
                if line:
                    parts.append(line)
        return ParsedDocument(
            filename=path.name,
            file_type=self.file_type,
            text="\n".join(parts),
            meta={"paragraphs": len(doc.paragraphs), "tables": len(doc.tables)},
        )
